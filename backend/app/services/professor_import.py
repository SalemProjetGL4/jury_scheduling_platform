from __future__ import annotations

import csv
import re
import unicodedata
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models


_COLUMN_ALIASES = {
    "professor_name": "name",
    "full_name": "name",
    "mail": "email",
    "department": "department_name",
    "faculty": "department_name",
    "jury_limit": "max_juries",
    "maxjury": "max_juries",
    "max_jury": "max_juries",
    "preferences_list": "preferences",
}


_REQUIRED_FIELDS = {"name", "email"}


def _normalize_header(value: str) -> str:
    normalized = value.strip().lower().replace("-", "_").replace(" ", "_")
    return _COLUMN_ALIASES.get(normalized, normalized)


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text: str
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.DictReader(StringIO(text))
    if reader.fieldnames is None:
        raise ValueError("CSV file has no header row")

    normalized_headers = [_normalize_header(header) for header in reader.fieldnames]
    rows: list[dict[str, Any]] = []

    for raw_row in reader:
        row: dict[str, Any] = {}
        for source_header, normalized in zip(reader.fieldnames, normalized_headers, strict=False):
            row[normalized] = raw_row.get(source_header)
        rows.append(row)

    return rows


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active

    raw_rows = list(worksheet.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("XLSX file is empty")

    first_row = raw_rows[0]
    headers = [
        _normalize_header(str(cell).strip()) if cell is not None else ""
        for cell in first_row
    ]

    if not any(headers):
        raise ValueError("XLSX file header row is empty")

    rows: list[dict[str, Any]] = []
    for values in raw_rows[1:]:
        if values is None:
            continue
        row: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = values[idx] if idx < len(values) else None
            row[header] = value
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            rows.append(row)

    return rows


def _load_rows(file_name: str, content: bytes) -> list[dict[str, Any]]:
    lowered = file_name.lower()
    if lowered.endswith(".csv"):
        return _parse_csv(content)
    if lowered.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise ValueError("Unsupported file format. Upload .csv or .xlsx")


def _to_int(value: Any) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _to_clean_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = normalized.lower()
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def _department_candidates(value: str) -> set[str]:
    normalized = _normalize_text(value)
    if not normalized:
        return set()

    compact = normalized.replace(" ", "")
    words = normalized.split()
    acronym = "".join(word[0] for word in words if word)

    candidates = {normalized, compact, acronym}
    for suffix in ("department", "dept", "faculty"):
        if normalized.endswith(f" {suffix}"):
            stripped = normalized[: -len(suffix) - 1].strip()
            candidates.add(stripped)
            candidates.add(stripped.replace(" ", ""))
            candidates.add("".join(word[0] for word in stripped.split() if word))
    return {candidate for candidate in candidates if candidate}


def _resolve_department_id(
    value: str | None,
    department_index: dict[str, int],
) -> int | None:
    if not value:
        return None

    for candidate in _department_candidates(value):
        resolved = department_index.get(candidate)
        if resolved is not None:
            return resolved
    return None


def _parse_preferences(value: Any) -> list[str] | None:
    text = _to_clean_str(value)
    if not text:
        return None
    return [item.strip() for item in text.replace(";", ",").split(",") if item.strip()]


def import_professors(
    *,
    db: Session,
    file_name: str,
    content: bytes,
    default_department_id: int | None,
    default_department_name: str | None,
    default_max_juries: int,
    dry_run: bool,
) -> dict[str, Any]:
    rows = _load_rows(file_name=file_name, content=content)

    if not rows:
        raise ValueError("Uploaded file has no professor rows")

    missing_required = [field for field in _REQUIRED_FIELDS if all(field not in row for row in rows)]
    if missing_required:
        raise ValueError(f"Missing required column(s): {', '.join(sorted(missing_required))}")

    departments = db.execute(select(models.Department)).scalars().all()
    department_ids = {department.id for department in departments}
    department_name_map: dict[str, int] = {}
    for department in departments:
        for candidate in _department_candidates(department.name):
            department_name_map.setdefault(candidate, department.id)

    if default_department_name:
        mapped_default = _resolve_department_id(default_department_name, department_name_map)
        if mapped_default is None:
            raise ValueError(f"Unknown default department_name: {default_department_name}")
        default_department_id = mapped_default

    existing_emails = {
        email.lower()
        for email in db.execute(select(models.Professor.email)).scalars().all()
        if email
    }
    professors_by_email = {
        professor.email.lower(): professor
        for professor in db.execute(select(models.Professor)).scalars().all()
        if professor.email
    }
    imported_emails: set[str] = set()

    issues: list[dict[str, Any]] = []
    professors_to_insert: list[models.Professor] = []
    updated_count = 0
    skipped_duplicates = 0

    for idx, row in enumerate(rows, start=2):
        name = _to_clean_str(row.get("name"))
        email_raw = _to_clean_str(row.get("email"))
        department_id = _to_int(row.get("department_id"))
        department_name = _to_clean_str(row.get("department_name"))
        max_juries = _to_int(row.get("max_juries")) or default_max_juries
        preferences = _parse_preferences(row.get("preferences"))

        if department_id is None and department_name:
            department_id = _resolve_department_id(department_name, department_name_map)
        if department_id is None and default_department_id is not None:
            department_id = default_department_id

        if not name:
            issues.append({"row_number": idx, "email": email_raw, "reason": "Missing name"})
            continue

        if not email_raw:
            issues.append({"row_number": idx, "email": None, "reason": "Missing email"})
            continue

        normalized_email = email_raw.lower()

        if department_id is None:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": "Missing department_id or department_name (provide a value or a default)",
                }
            )
            continue

        if department_id not in department_ids:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": f"Unknown department_id: {department_id}",
                }
            )
            continue

        if max_juries is None:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": "Missing max_juries",
                }
            )
            continue

        if normalized_email in imported_emails:
            skipped_duplicates += 1
            continue

        imported_emails.add(normalized_email)
        existing_professor = professors_by_email.get(normalized_email)
        if existing_professor is not None:
            existing_professor.name = name
            existing_professor.department_id = department_id
            existing_professor.max_juries = max_juries
            existing_professor.preferences = preferences
            updated_count += 1
            continue

        professors_to_insert.append(
            models.Professor(
                name=name,
                email=normalized_email,
                department_id=department_id,
                max_juries=max_juries,
                preferences=preferences,
            )
        )

    if not dry_run and professors_to_insert:
        db.add_all(professors_to_insert)
        db.commit()

    return {
        "file_name": file_name,
        "dry_run": dry_run,
        "total_rows": len(rows),
        "inserted_count": len(professors_to_insert),
        "updated_count": updated_count,
        "skipped_duplicates": skipped_duplicates,
        "invalid_rows": len(issues),
        "issues": issues,
    }