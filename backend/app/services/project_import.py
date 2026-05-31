from __future__ import annotations

import csv
import re
import unicodedata
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession

from app import models
from app.services.domain_mapping import infer_domain_ids


_HEADER_MAP: dict[str, str] = {
    # student_name
    "student_name": "student_name",
    "nom_etudiant": "student_name",
    "student": "student_name",
    "etudiant": "student_name",
    # student_email
    "student_email": "student_email",
    "email_etudiant": "student_email",
    "email": "student_email",
    # filiere
    "filiere": "filiere",
    "filière": "filiere",
    "degree": "filiere",
    "programme": "filiere",
    # project_title
    "project_title": "project_title",
    "title": "project_title",
    "titre": "project_title",
    "projet": "project_title",
    # domain
    "domain": "domain",
    "domaine": "domain",
    "domain_name": "domain",
    # supervisor_name
    "supervisor_name": "supervisor_name",
    "supervisor": "supervisor_name",
    "encadrant": "supervisor_name",
    "nom_encadrant": "supervisor_name",
    # enterprise
    "enterprise": "enterprise",
    "entreprise": "enterprise",
    "company": "enterprise",
    # enterprise_supervisor
    "enterprise_supervisor": "enterprise_supervisor",
    "encadrant_entreprise": "enterprise_supervisor",
    "enterprise_supervisor_name": "enterprise_supervisor",
    "tuteur_entreprise": "enterprise_supervisor",
}

_REQUIRED_KEYS = {
    "student_name",
    "student_email",
    "filiere",
    "project_title",
    "domain",
    "supervisor_name",
}


def normalize_prof_name(name: str) -> str:
    name = name.strip()
    name = unicodedata.normalize("NFD", name)
    name = "".join(c for c in name if unicodedata.category(c) != "Mn")
    name = name.lower()
    name = re.sub(r"[^a-z0-9]+", " ", name)
    name = " ".join(name.split())
    return name


def _normalize(value: str) -> str:
    return value.strip().lower().replace(" ", "_").replace("-", "_")


def _map_header(raw: str) -> str:
    key = _normalize(raw)
    return _HEADER_MAP.get(key, key)


def _detect_delimiter(text: str) -> str:
    first_line = text.split("\n", 1)[0]
    return ";" if first_line.count(";") >= first_line.count(",") else ","


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    delimiter = _detect_delimiter(text)
    reader = csv.DictReader(StringIO(text), delimiter=delimiter)
    if reader.fieldnames is None:
        raise ValueError("CSV file has no header row")

    mapped = [_map_header(h) for h in reader.fieldnames]
    rows: list[dict[str, Any]] = []
    for raw_row in reader:
        row: dict[str, Any] = {}
        for src, dst in zip(reader.fieldnames, mapped):
            row[dst] = raw_row.get(src)
        rows.append(row)
    return rows


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("XLSX file is empty")

    headers = [
        _map_header(str(cell).strip()) if cell is not None else ""
        for cell in raw_rows[0]
    ]
    if not any(headers):
        raise ValueError("XLSX header row is empty")

    rows: list[dict[str, Any]] = []
    for values in raw_rows[1:]:
        if values is None:
            continue
        row: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            row[header] = values[idx] if idx < len(values) else None
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            rows.append(row)
    return rows


def _load_rows(file_name: str, content: bytes) -> list[dict[str, Any]]:
    if file_name.lower().endswith(".csv"):
        return _parse_csv(content)
    if file_name.lower().endswith(".xlsx"):
        return _parse_xlsx(content)
    raise ValueError("Unsupported file format. Upload .csv or .xlsx")


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _resolve_domain_ids(
    raw_names: str | None,
    raw_ids: str | None,
    domain_name_map: dict[str, int],
    all_domain_ids: set[int],
    default_domain_id: int | None,
) -> tuple[list[int], str | None]:
    """Return (resolved_ids, error_message). Supports comma-separated names/ids."""
    resolved: list[int] = []

    # Parse comma-separated IDs from a domain_id column if present
    if raw_ids:
        for part in str(raw_ids).split(","):
            did = _to_int(part.strip())
            if did is not None:
                resolved.append(did)

    # Parse comma-separated names from domain_name/domain_names column
    if raw_names:
        for part in str(raw_names).split(","):
            name = part.strip().lower()
            if name:
                did = domain_name_map.get(name)
                if did is not None and did not in resolved:
                    resolved.append(did)

    # Fall back to default
    if not resolved and default_domain_id is not None:
        resolved.append(default_domain_id)

    if not resolved:
        return [], "Missing domain_id or domain_name"

    unknown = [d for d in resolved if d not in all_domain_ids]
    if unknown:
        return [], f"Unknown domain id(s): {unknown}"

    return resolved, None


def import_projects(
    *,
    db: DBSession,
    file_name: str,
    content: bytes,
    session_id: int,
    dry_run: bool = False,
) -> dict[str, Any]:
    rows = _load_rows(file_name=file_name, content=content)
    if not rows:
        return {"status": "error", "message": "No data rows found in file", "created": 0, "skipped": 0, "issues": []}

    # Validate required columns are present in the file
    file_keys: set[str] = set()
    for row in rows:
        file_keys.update(row.keys())
    missing = sorted(_REQUIRED_KEYS - file_keys)
    if missing:
        return {
            "status": "error",
            "message": f"Missing required columns: {missing}",
            "created": 0,
            "skipped": 0,
            "issues": [],
        }

    # Verify session exists and get its start_date for new student promotions
    session = db.get(models.Session, session_id)
    if session is None:
        return {
            "status": "error",
            "message": f"Session {session_id} not found",
            "created": 0,
            "skipped": 0,
            "issues": [],
        }

    # Pre-load lookup tables
    domains_by_name: dict[str, int] = {
        d.name.strip().lower(): d.id
        for d in db.execute(select(models.Domain)).scalars().all()
    }
    filieres_by_name: dict[str, int] = {
        f.name.strip().lower(): f.id
        for f in db.execute(select(models.Filiere)).scalars().all()
    }
    professors_by_name: dict[str, int] = {
        normalize_prof_name(p.name): p.id
        for p in db.execute(select(models.Professor)).scalars().all()
    }
    existing_emails: set[str] = {
        e.lower()
        for e in db.execute(select(models.Student.email)).scalars().all()
    }

    issues: list[str] = []
    created = 0
    skipped = 0

    for row_num, row in enumerate(rows, start=2):
        student_name = _clean(row.get("student_name"))
        student_email = _clean(row.get("student_email"))
        filiere_raw = _clean(row.get("filiere"))
        project_title = _clean(row.get("project_title"))
        domain_raw = _clean(row.get("domain"))
        supervisor_raw = _clean(row.get("supervisor_name"))
        enterprise = _clean(row.get("enterprise"))
        enterprise_supervisor = _clean(row.get("enterprise_supervisor"))

        # Step 1 — Resolve supervisor
        if not supervisor_raw:
            issues.append(f"row {row_num}: Missing supervisor name — row skipped")
            skipped += 1
            continue
        supervisor_id = professors_by_name.get(normalize_prof_name(supervisor_raw))
        if supervisor_id is None:
            issues.append(f"row {row_num}: Supervisor '{supervisor_raw}' not found in database — row skipped")
            skipped += 1
            continue

        # Step 2 — Resolve domain
        if not domain_raw:
            issues.append(f"row {row_num}: Missing domain — row skipped")
            skipped += 1
            continue
        domain_id = domains_by_name.get(domain_raw.lower())
        if domain_id is None:
            issues.append(f"row {row_num}: Domain '{domain_raw}' not found — row skipped")
            skipped += 1
            continue

        # Step 3 — Resolve filiere
        if not filiere_raw:
            issues.append(f"row {row_num}: Missing filière — row skipped")
            skipped += 1
            continue
        filiere_id = filieres_by_name.get(filiere_raw.lower())
        if filiere_id is None:
            issues.append(f"row {row_num}: Filière '{filiere_raw}' not found — row skipped")
            skipped += 1
            continue

        # Validate remaining required fields
        if not student_email:
            issues.append(f"row {row_num}: Missing student email — row skipped")
            skipped += 1
            continue
        if not student_name:
            issues.append(f"row {row_num}: Missing student name — row skipped")
            skipped += 1
            continue
        if not project_title:
            issues.append(f"row {row_num}: Missing project title — row skipped")
            skipped += 1
            continue

        # Step 4 — Upsert student
        if student_email.lower() in existing_emails:
            student = db.execute(
                select(models.Student).where(models.Student.email.ilike(student_email))
            ).scalar_one_or_none()
            if student is not None and student.name != student_name:
                student.name = student_name
        else:
            student = models.Student(
                name=student_name,
                email=student_email,
                filiere_id=filiere_id,
                promotion=session.start_date,
            )
            db.add(student)
            db.flush()
            existing_emails.add(student_email.lower())

        if student is None:
            issues.append(f"row {row_num}: Could not create or find student '{student_email}' — row skipped")
            skipped += 1
            continue

        # Step 5 — Create project
        db.add(models.Project(
            title=project_title,
            domain_id=domain_id,
            supervisor_id=supervisor_id,
            student_id=student.id,
            session_id=session_id,
            enterprise=enterprise,
            enterprise_supervisor=enterprise_supervisor,
        ))
        created += 1

    # Step 6 — Commit or rollback
    if dry_run:
        db.rollback()
        return {"status": "dry_run", "created": created, "skipped": skipped, "issues": issues}

    db.commit()
    if created == 0:
        status = "error"
    elif skipped > 0:
        status = "partial"
    else:
        status = "ok"

    return {"status": status, "created": created, "skipped": skipped, "issues": issues}
