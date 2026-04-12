from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from pydantic import EmailStr, TypeAdapter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import models


_EMAIL_ADAPTER = TypeAdapter(EmailStr)

_COLUMN_ALIASES = {
    "student_name": "name",
    "full_name": "name",
    "mail": "email",
    "academic_year": "promotion_year",
    "year": "promotion_year",
    "filiere": "filiere_name",
    "track": "filiere_name",
    "department": "filiere_name",
}

_REQUIRED_FIELDS = {"name", "email"}


def _normalize_header(value: str) -> str:
    normalized = value.strip().lower().replace("-", "_").replace(" ", "_")
    return _COLUMN_ALIASES.get(normalized, normalized)


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text: str
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    reader = csv.DictReader(StringIO(text))
    if reader.fieldnames is None:
        raise ValueError("CSV file has no header row")

    normalized_headers = [_normalize_header(header) for header in reader.fieldnames]
    rows: list[dict[str, Any]] = []

    for raw_row in reader:
        row: dict[str, Any] = {}
        for source_header, normalized in zip(reader.fieldnames, normalized_headers, strict=False):
            row[normalized] = raw_row.get(source_header)
        rows.append(row)

    return rows


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active

    raw_rows = list(worksheet.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("XLSX file is empty")

    first_row = raw_rows[0]
    headers = [
        _normalize_header(str(cell).strip()) if cell is not None else ""
        for cell in first_row
    ]

    if not any(headers):
        raise ValueError("XLSX file header row is empty")

    rows: list[dict[str, Any]] = []
    for values in raw_rows[1:]:
        if values is None:
            continue
        row: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = values[idx] if idx < len(values) else None
            row[header] = value
        if any(v is not None and str(v).strip() != "" for v in row.values()):
            rows.append(row)

    return rows


def _load_rows(file_name: str, content: bytes) -> list[dict[str, Any]]:
    lowered = file_name.lower()
    if lowered.endswith(".csv"):
        return _parse_csv(content)
    if lowered.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise ValueError("Unsupported file format. Upload .csv or .xlsx")


def _to_int(value: Any) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _to_clean_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_students(
    *,
    db: Session,
    file_name: str,
    content: bytes,
    default_promotion_year: int | None,
    default_filiere_id: int | None,
    default_filiere_name: str | None,
    dry_run: bool,
) -> dict[str, Any]:
    rows = _load_rows(file_name=file_name, content=content)

    if not rows:
        raise ValueError("Uploaded file has no student rows")

    missing_required = [field for field in _REQUIRED_FIELDS if all(field not in row for row in rows)]
    if missing_required:
        raise ValueError(f"Missing required column(s): {', '.join(sorted(missing_required))}")

    filieres = db.execute(select(models.Filiere)).scalars().all()
    filiere_ids = {f.id for f in filieres}
    filiere_name_map = {f.name.strip().lower(): f.id for f in filieres}

    if default_filiere_name:
        mapped_default = filiere_name_map.get(default_filiere_name.strip().lower())
        if mapped_default is None:
            raise ValueError(f"Unknown default filiere_name: {default_filiere_name}")
        default_filiere_id = mapped_default

    existing_emails = {
        email.lower()
        for email in db.execute(select(models.Student.email)).scalars().all()
        if email
    }
    imported_emails: set[str] = set()

    issues: list[dict[str, Any]] = []
    students_to_insert: list[models.Student] = []
    skipped_duplicates = 0

    for idx, row in enumerate(rows, start=2):
        name = _to_clean_str(row.get("name"))
        email_raw = _to_clean_str(row.get("email"))
        promotion_year = _to_int(row.get("promotion_year")) or default_promotion_year

        row_filiere_id = _to_int(row.get("filiere_id"))
        row_filiere_name = _to_clean_str(row.get("filiere_name"))

        if row_filiere_id is None and row_filiere_name:
            row_filiere_id = filiere_name_map.get(row_filiere_name.lower())

        filiere_id = row_filiere_id or default_filiere_id

        if not name:
            issues.append({"row_number": idx, "email": email_raw, "reason": "Missing name"})
            continue

        if not email_raw:
            issues.append({"row_number": idx, "email": None, "reason": "Missing email"})
            continue

        try:
            normalized_email = str(_EMAIL_ADAPTER.validate_python(email_raw)).lower()
        except Exception:
            issues.append({"row_number": idx, "email": email_raw, "reason": "Invalid email format"})
            continue

        if promotion_year is None:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": "Missing promotion_year and no default provided",
                }
            )
            continue

        if filiere_id is None:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": "Missing filiere (provide filiere_id/filiere_name or a default)",
                }
            )
            continue

        if filiere_id not in filiere_ids:
            issues.append(
                {
                    "row_number": idx,
                    "email": normalized_email,
                    "reason": f"Unknown filiere_id: {filiere_id}",
                }
            )
            continue

        if normalized_email in existing_emails or normalized_email in imported_emails:
            skipped_duplicates += 1
            continue

        imported_emails.add(normalized_email)
        students_to_insert.append(
            models.Student(
                name=name,
                email=normalized_email,
                promotion_year=promotion_year,
                filiere_id=filiere_id,
            )
        )

    if not dry_run and students_to_insert:
        db.add_all(students_to_insert)
        db.commit()

    return {
        "file_name": file_name,
        "dry_run": dry_run,
        "total_rows": len(rows),
        "inserted_count": len(students_to_insert),
        "skipped_duplicates": skipped_duplicates,
        "invalid_rows": len(issues),
        "issues": issues,
    }
